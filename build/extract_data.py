"""
Step 1 of the build pipeline: spreadsheet -> JSON.

Reads the `tblSalesTransactions` Excel Table on the "Sales" sheet of
data/Sales Data.xlsx and writes a flat JSON array to data/sales.json.

This is the ONE file you'd change to point the dashboard at a database
instead of a spreadsheet later (see documentation/Documentation.md, "Moving
to a database"): swap the openpyxl read below for a DB query that yields the
same list-of-dicts shape, keep everything from build_dashboard.py onwards
unchanged.

Run from inside build/:
    python3 extract_data.py
"""

import json
from datetime import date, datetime

import openpyxl

SOURCE_XLSX = "../data/Sales Data.xlsx"
SHEET_NAME = "Sales"
TABLE_NAME = "tblSalesTransactions"
OUTPUT_JSON = "../data/sales.json"


def find_table_range(ws, table_name):
    tbl = ws.tables.get(table_name)
    if tbl is None:
        raise ValueError(
            f"Table '{table_name}' not found on sheet '{ws.title}'. "
            f"Available tables: {list(ws.tables.keys())}"
        )
    return tbl.ref  # e.g. "A1:M7494"


def to_iso_date(value):
    """None/blank stays None (caller decides the fallback); anything else
    is coerced to a 'yyyy-mm-dd' string."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)[:10]


def to_float(value, default=0.0):
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return default


def to_int(value, default=0):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def main():
    wb = openpyxl.load_workbook(SOURCE_XLSX, data_only=True)
    ws = wb[SHEET_NAME]

    table_ref = find_table_range(ws, TABLE_NAME)
    cell_range = ws[table_ref]

    header = [c.value for c in cell_range[0]]
    records = []
    skipped_no_date = 0
    for row in cell_range[1:]:
        values = [c.value for c in row]
        if all(v is None for v in values):
            continue
        rec = dict(zip(header, values))

        if not rec.get("OrderID"):
            continue

        # SaleDate drives year/month/monthLabel grouping (charts, the Year
        # filter, the pivot table) — PurchaseDate is carried through as a
        # separate field but isn't used for any time-based grouping. Real
        # spreadsheets are often incomplete on one of the two dates, so fall
        # back to whichever is present rather than dropping the row; only
        # skip it if NEITHER date is available (nothing to group it by).
        iso_purchase_date = to_iso_date(rec.get("PurchaseDate"))
        iso_sale_date = to_iso_date(rec.get("SaleDate"))
        if iso_sale_date is None:
            iso_sale_date = iso_purchase_date
        if iso_purchase_date is None:
            iso_purchase_date = iso_sale_date
        if iso_sale_date is None:
            skipped_no_date += 1
            continue

        y, m, _ = iso_sale_date.split("-")
        month_label = datetime.strptime(f"{y}-{m}", "%Y-%m").strftime("%b %Y")

        records.append({
            "orderId": rec["OrderID"],
            "purchaseDate": iso_purchase_date,
            "saleDate": iso_sale_date,
            "year": int(y),
            "month": int(m),
            "monthLabel": month_label,
            # Blank category/location/segment/salesperson are bucketed under
            # a labelled placeholder rather than left null — a null value
            # can't be filtered on or coloured consistently downstream.
            "category": rec.get("Category") or "Uncategorized",
            "product": rec.get("Product") or "Unspecified",
            "saleLocation": rec.get("SaleLocation") or "Unspecified",
            "segment": rec.get("CustomerSegment") or "Unspecified",
            "salesperson": rec.get("Salesperson") or "Unassigned",
            "units": to_int(rec.get("Units"), default=1),
            "unitPrice": to_float(rec.get("UnitPrice")),
            "revenue": to_float(rec.get("Revenue")),
            "cost": to_float(rec.get("Cost")),
            "profit": to_float(rec.get("Profit")),
        })

    with open(OUTPUT_JSON, "w") as f:
        json.dump(records, f, separators=(",", ":"))

    print(f"Extracted {len(records):,} rows from '{TABLE_NAME}' "
          f"-> {OUTPUT_JSON}"
          + (f" ({skipped_no_date} row(s) skipped: no PurchaseDate or "
             f"SaleDate)" if skipped_no_date else ""))


if __name__ == "__main__":
    main()
